"""Document parsing (FR-015): bytes + content-type -> plain text.

Text/Markdown/HTML/CSV/TSV/JSON are handled with the stdlib. PDF/DOCX/XLSX use
optional libraries from the `data` extra, imported lazily so the core has no
heavy dependencies.
"""

from __future__ import annotations

import csv
import io
import os
from html.parser import HTMLParser

from eadip.ingestion.models import ParsedDocument


class UnsupportedContentType(ValueError):
    pass


_EXT_TO_CT = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".csv": "text/csv",
    ".tsv": "text/tab-separated-values",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def guess_content_type(filename: str) -> str:
    return _EXT_TO_CT.get(os.path.splitext(filename)[1].lower(), "application/octet-stream")


class _HTMLTextExtractor(HTMLParser):
    _SKIP = {"script", "style", "head"}

    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []
        self._skip_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in self._SKIP:
            self._skip_depth += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1

    def handle_data(self, data: str) -> None:
        if not self._skip_depth and data.strip():
            self._parts.append(data.strip())

    def text(self) -> str:
        return " ".join(self._parts)


def _norm(content_type: str) -> str:
    return content_type.split(";")[0].strip().lower()


def _decode(raw: bytes) -> str:
    return raw.decode("utf-8", errors="replace")


def parse(raw: bytes, content_type: str, *, filename: str | None = None) -> ParsedDocument:
    ct = _norm(content_type)
    if ct in ("application/octet-stream", "") and filename:
        ct = _norm(guess_content_type(filename))

    if ct in ("text/plain", "text/markdown", "text/x-markdown", "application/json"):
        return ParsedDocument(text=_decode(raw), content_type=ct)
    if ct in ("text/html", "application/xhtml+xml"):
        parser = _HTMLTextExtractor()
        parser.feed(_decode(raw))
        return ParsedDocument(text=parser.text(), content_type=ct)
    if ct in ("text/csv", "text/tab-separated-values"):
        delimiter = "\t" if ct.endswith("tab-separated-values") else ","
        rows = csv.reader(io.StringIO(_decode(raw)), delimiter=delimiter)
        text = "\n".join(" ".join(cell.strip() for cell in row) for row in rows)
        return ParsedDocument(text=text, content_type=ct)
    if ct == "application/pdf":
        return ParsedDocument(text=_parse_pdf(raw), content_type=ct)
    if ct == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return ParsedDocument(text=_parse_docx(raw), content_type=ct)
    if ct == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return ParsedDocument(text=_parse_xlsx(raw), content_type=ct)
    raise UnsupportedContentType(f"unsupported content type: {content_type}")


def _parse_pdf(raw: bytes) -> str:
    try:
        import pypdf
    except ImportError as exc:  # pragma: no cover - exercised only without extra
        raise UnsupportedContentType("install the 'data' extra for PDF support") from exc
    reader = pypdf.PdfReader(io.BytesIO(raw))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _parse_docx(raw: bytes) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedContentType("install the 'data' extra for DOCX support") from exc
    document = docx.Document(io.BytesIO(raw))
    return "\n".join(p.text for p in document.paragraphs)


def _parse_xlsx(raw: bytes) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise UnsupportedContentType("install the 'data' extra for XLSX support") from exc
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            lines.append(" ".join("" if cell is None else str(cell) for cell in row))
    return "\n".join(lines)
